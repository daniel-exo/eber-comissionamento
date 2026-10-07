"""Gera a lista de ações de inspeção do app (dados/acoes.json) a partir da planilha de ações.

Uso:  python ferramentas/gerar_acoes.py <planilha.xlsx> [conferencia.csv] [--simular]

A planilha precisa de uma aba com as colunas "Tag" e "Problema" (e, se houver, "Data adicionado" e "Nº").
  - Tag: o local completo do equipamento na lista (ex.: 400.DCT-4001.FCV-40001) ou só a área (ex.: 400),
    quando a ação é geral da área.
  - Nº: número próprio da ação. Nunca muda e nunca é reaproveitado. É por ele que o check fica guardado no banco.

Números (o que mantém a base):
  - Linha com Nº na planilha: vale o número da planilha. O gerador confere com o registro
    (ferramentas/registro-acoes.json); se o número já existia com outro tag E outro texto, para sem gravar nada,
    porque isso indica número reaproveitado ou renumeração.
  - Linha sem Nº: o gerador procura a mesma ação no registro (mesmo tag e mesmo texto) e mantém o número dela;
    se não achar, é ação nova e recebe o próximo número livre.
  - Ação que sai da planilha some do app, mas o número fica reservado e o check continua no banco.

Ficam fora do app (com o número reservado): ações de equipamentos listados em ferramentas/expansao.txt
e ações de áreas que o app não tem.

Nada é gravado quando há erro, nem com --simular. O banco nunca é tocado por este script.
"""
import csv, datetime, difflib, json, os, re, sys, unicodedata
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
REG = os.path.join(AQUI, 'registro-acoes.json')
SAIDA = os.path.join(RAIZ, 'dados', 'acoes.json')
EXPANSAO = os.path.join(AQUI, 'expansao.txt')
AREAS = ['200', '300', '400', '700']


def norm(s):
    """Texto sem acento, minúsculo e com espaços simples — só para comparar."""
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', s).strip().casefold()


def limpo(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def ler_expansao():
    if not os.path.exists(EXPANSAO):
        return []
    return [l.split('#', 1)[0].strip().upper() for l in open(EXPANSAO, encoding='utf-8') if l.split('#', 1)[0].strip()]


def na_expansao(loc, exclusoes):
    segs = loc.split('.')[1:]
    return any(('.' in e and (loc == e or loc.startswith(e + '.'))) or ('.' not in e and e in segs) for e in exclusoes)


def ler_planilha(xlsx):
    wb = openpyxl.load_workbook(xlsx)          # sem data_only: uma fórmula na coluna Nº aparece como texto "=..."
    achou = None
    for ws in wb.worksheets:
        cab = [norm(c.value) for c in ws[1]]
        if 'tag' in cab and any(x in cab for x in ('problema', 'acao')):
            achou = (ws, cab)
            break
    if not achou:
        raise SystemExit('ERRO: não encontrei uma aba com as colunas "Tag" e "Problema".')
    ws, cab = achou
    col = lambda *nomes: next((cab.index(n) for n in nomes if n in cab), None)
    i_tag, i_txt = col('tag'), col('problema', 'acao')
    i_data = next((i for i, c in enumerate(cab) if c.startswith('data')), None)
    i_num = col('nº', 'n°', 'no', 'n', 'numero', 'num', '#', 'id')
    linhas = []
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, '') for v in r):
            continue
        d = r[i_data] if i_data is not None else None
        d = d.date().isoformat() if isinstance(d, datetime.datetime) else d.isoformat() if isinstance(d, datetime.date) else ''
        tag = r[i_tag]
        tag = str(int(tag)) if isinstance(tag, (int, float)) and float(tag).is_integer() else limpo(tag)
        linhas.append({'linha': n, 'num': r[i_num] if i_num is not None else None, 'tag': tag.upper(), 'texto': limpo(r[i_txt]), 'data': d})
    return ws.title, i_num is not None, linhas


def lista_do_app():
    """local -> posição na lista do app (para ordenar e para saber se o tag existe)."""
    pos = {}
    for a in AREAS:
        d = json.load(open(os.path.join(RAIZ, 'dados', f'area-{a}.json'), encoding='utf-8'))
        for c in d['centrais']:
            for it in c['itens']:
                pos[it['loc']] = len(pos)
    return pos


def main(xlsx, conferencia=None, simular=False):
    aba, tem_coluna, linhas = ler_planilha(xlsx)
    reg = json.load(open(REG, encoding='utf-8')) if os.path.exists(REG) else {'proximo': 1, 'acoes': {}}
    antigas = {int(k): v for k, v in reg['acoes'].items()}
    primeira_vez = not antigas
    erros, avisos = [], []

    # ---- linhas inválidas
    validas = []
    for l in linhas:
        if not l['tag']:
            avisos.append(f"linha {l['linha']}: sem Tag — ignorada"); continue
        if not l['texto']:
            avisos.append(f"linha {l['linha']}: sem texto do problema — ignorada"); continue
        validas.append(l)

    # ---- 1) linhas com Nº na planilha
    usados, por_num = set(), {}
    for l in validas:
        v = l['num']
        if v in (None, ''):
            l['n'] = None; continue
        if isinstance(v, str) and v.strip().startswith('='):
            erros.append(f"linha {l['linha']}: a coluna Nº tem fórmula; o número precisa ser digitado (valor fixo)"); l['n'] = None; continue
        try:
            n = int(str(v).strip().upper().replace('AC-', ''))
            assert n > 0 and float(str(v).strip().upper().replace('AC-', '')) == n
        except Exception:
            erros.append(f"linha {l['linha']}: Nº inválido ({v!r})"); l['n'] = None; continue
        if n in por_num:
            erros.append(f"Nº {n} repetido nas linhas {por_num[n]} e {l['linha']}")
        por_num[n] = l['linha']
        l['n'] = n
        usados.add(n)
        if n in antigas:
            a = antigas[n]
            if norm(a['tag']) != norm(l['tag']) and norm(a['texto']) != norm(l['texto']):
                erros.append(f"Nº {n} (linha {l['linha']}): no registro é \"{a['tag']} — {a['texto']}\" e na planilha é "
                             f"\"{l['tag']} — {l['texto']}\". Número reaproveitado ou planilha renumerada?")

    # ---- 2) linhas sem Nº: procura a mesma ação no registro
    livres = {n: a for n, a in antigas.items() if n not in usados}
    def pega(l, chave):
        cand = [n for n, a in livres.items() if chave(a) == chave(l)]
        if len(cand) == 1 or (len(cand) > 1 and chave is ch_completa):
            n = min(cand); del livres[n]; return n
        return None
    ch_completa = lambda x: (norm(x['tag']), norm(x['texto']), x.get('data') or '')
    ch_sem_data = lambda x: (norm(x['tag']), norm(x['texto']))
    sem = [l for l in validas if l.get('n') is None and not any(e.startswith(f"linha {l['linha']}:") for e in erros)]
    for chave in (ch_completa, ch_sem_data):
        for l in sem:
            if l['n'] is None:
                l['n'] = pega(l, chave)
    # texto corrigido numa linha sem número: mesmo tag, mesma data e texto muito parecido, sem ambiguidade
    parecidas = []
    for l in [x for x in sem if x['n'] is None]:
        cand = [n for n, a in livres.items() if not a.get('ausente') and norm(a['tag']) == norm(l['tag']) and (a.get('data') or '') == l['data']
                and difflib.SequenceMatcher(None, norm(a['texto']), norm(l['texto'])).ratio() >= 0.8]
        outras = [x for x in sem if x['n'] is None and x is not l and norm(x['tag']) == norm(l['tag']) and x['data'] == l['data']]
        if len(cand) == 1 and not outras:
            l['n'] = cand[0]; parecidas.append(l); del livres[cand[0]]
    # ---- 3) o que sobrou é ação nova
    prox = max([reg.get('proximo', 1)] + [n + 1 for n in list(antigas) + list(usados)])
    novas = []
    for l in sem:
        if l['n'] is None:
            l['n'] = prox; prox += 1; novas.append(l)
    novas += [l for l in validas if l.get('n') in usados and l['n'] not in antigas]
    novas.sort(key=lambda l: l['n'])

    if erros:
        print('NADA FOI GRAVADO. Corrija a planilha:'); [print('  -', e) for e in erros]
        return 1

    # ---- situação de cada ação
    pos, exclusoes = lista_do_app(), ler_expansao()
    alteradas = []
    for l in validas:
        area = l['tag'].split('.')[0]
        l['area'] = area
        geral = '.' not in l['tag']
        l['loc'] = '' if geral else l['tag']
        if area not in AREAS: l['sit'] = f'fora do app (área {area} não está no app)'
        elif geral: l['sit'] = 'app'
        elif na_expansao(l['tag'], exclusoes): l['sit'] = 'fora do app (equipamento da expansão)'
        elif l['tag'] in pos: l['sit'] = 'app'
        else: l['sit'] = 'app (tag não encontrado na lista de equipamentos)'
        a = antigas.get(l['n'])
        if a and (norm(a['tag']) != norm(l['tag']) or a['texto'] != l['texto']) and l not in novas:
            alteradas.append((l, a))
    presentes = {l['n'] for l in validas}
    sairam = [(n, a) for n, a in sorted(antigas.items()) if n not in presentes and not a.get('ausente')]

    # ---- arquivo do app
    ident = lambda n: f'AC-{n:04d}'
    areas = {a: [] for a in AREAS}
    for l in validas:
        if not l['sit'].startswith('app'):
            continue
        segs = l['tag'].split('.')[1:]
        rot = lambda s: f'Grupo {s}' if re.fullmatch(r'[A-D]', s) else s
        item = {'id': ident(l['n']), 'tag': segs[-1] if segs else '', 'pai': ' › '.join(rot(s) for s in segs[:-1]), 't': l['texto'], 'd': l['data']}
        if l['loc']: item['loc'] = l['loc']
        if 'não encontrado' in l['sit']: item['nl'] = 1
        ordem = (0, 0, '') if not l['loc'] else (1, pos[l['loc']], '') if l['loc'] in pos else (2, 0, l['loc'])
        areas[l['area']].append((ordem + (l['n'],), item))
    saida = {'v': 1, 'gerado': datetime.datetime.now().isoformat(timespec='seconds'),
             'areas': {a: [it for _, it in sorted(v, key=lambda x: x[0])] for a, v in areas.items()}}

    # ---- registro
    novo_reg = {}
    for n, a in antigas.items():
        novo_reg[str(n)] = dict(a, ausente=True) if n not in presentes else a
    for l in validas:
        novo_reg[str(l['n'])] = {'tag': l['tag'], 'texto': l['texto'], 'data': l['data'], 'situacao': l['sit']}

    # ---- relatório
    p = print
    p(f'Planilha: aba "{aba}", {len(linhas)} linhas com conteúdo, {len(validas)} ações válidas.')
    p(f'Coluna Nº na planilha: {"sim" if tem_coluna else "NÃO"}'
      + (f' ({sum(1 for l in validas if l["num"] not in (None, ""))} linhas preenchidas)' if tem_coluna else ' — números atribuídos pela ordem das linhas e guardados no registro'))
    p('Ações no app por área: ' + ' · '.join(f'{a}: {len(saida["areas"][a])}' for a in AREAS) + f'  (total {sum(len(v) for v in saida["areas"].values())})')
    p(f'Números: {min(presentes)} a {max(presentes)}; próximo livre: {prox}')
    if primeira_vez:
        p('Primeira geração: todas as ações são novas.')
    else:
        p(f'Novas: {len(novas)}'); [p(f'   {ident(l["n"])}  {l["tag"]} — {l["texto"]}') for l in novas]
        p(f'Alteradas (mesmo número): {len(alteradas)}')
        for l, a in alteradas:
            p(f'   {ident(l["n"])}  antes: {a["tag"]} — {a["texto"]}'); p(f'            agora: {l["tag"]} — {l["texto"]}')
        if parecidas:
            p(f'   das quais {len(parecidas)} reconhecida(s) por semelhança (linha sem Nº com texto corrigido): ' + ', '.join(ident(l['n']) for l in parecidas))
        p(f'Saíram da planilha (somem do app; número reservado, check mantido no banco): {len(sairam)}')
        [p(f'   {ident(n)}  {a["tag"]} — {a["texto"]}') for n, a in sairam]
    fora = [l for l in validas if l['sit'].startswith('fora')]
    p(f'Fora do app, com número reservado: {len(fora)}'); [p(f'   {ident(l["n"])}  linha {l["linha"]}: {l["tag"]} — {l["sit"]}') for l in fora]
    nl = [l for l in validas if 'não encontrado' in l['sit']]
    p(f'Tag não encontrado na lista (a ação entra no app mesmo assim; confira a digitação): {len(nl)}'); [p(f'   {ident(l["n"])}  linha {l["linha"]}: {l["tag"]}') for l in nl]
    if avisos:
        p('Avisos:'); [p('   ' + a) for a in avisos]

    if conferencia:
        with open(conferencia, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(['Nº', 'Código no app', 'Linha da planilha', 'Área', 'Tag', 'Problema', 'Data adicionado', 'Situação'])
            for l in sorted(validas, key=lambda x: x['linha']):
                w.writerow([l['n'], ident(l['n']), l['linha'], l['area'], l['tag'], l['texto'], l['data'], l['sit']])
    if simular:
        p('\nSIMULAÇÃO: nada foi gravado.')
        return 0
    json.dump(saida, open(SAIDA, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    json.dump({'proximo': prox, 'acoes': dict(sorted(novo_reg.items(), key=lambda kv: int(kv[0])))}, open(REG, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    p('\nGravado: dados/acoes.json e ferramentas/registro-acoes.json')
    return 0


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        raise SystemExit(__doc__)
    sys.exit(main(args[0], args[1] if len(args) > 1 else None, '--simular' in sys.argv))
